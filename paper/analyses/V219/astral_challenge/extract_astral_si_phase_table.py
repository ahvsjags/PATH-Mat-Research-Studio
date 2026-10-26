#!/usr/bin/env python3
"""Extract the structured SI3.5 ASTRAL phase-fraction table and fail closed.

This produces a provenance-linked partial transcription for audit only. It does
not call the 202 rows a complete 224-reaction cohort, train PATH-Mat, or compute
external predictive performance. The published count mismatch must be resolved
from the publisher's full Supplementary Data 2 before any confirmatory scoring.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import fitz
import pandas as pd
from openpyxl import load_workbook


TEMP_RE = re.compile(r"^(600|700|800|900|1000)$")
FLOAT_RE = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def extract(pdf_path: Path) -> tuple[pd.DataFrame, dict]:
    doc = fitz.open(pdf_path)
    rows: list[dict] = []
    page_audit = []
    for page_index in range(32, 67):  # PDF pages 33–67, SI3.5 target tables.
        page = doc[page_index]
        lines = [s.strip() for s in page.get_text().splitlines()]
        table_start = next((i for i, s in enumerate(lines)
                            if s.startswith("Table 1. Temperature, XRD intensities")), None)
        if table_start is None:
            page_audit.append({"page": page_index + 1, "status": "missing_table_1"})
            continue
        table_end = next((i for i in range(table_start + 1, len(lines))
                          if lines[i].startswith("Table 2.")), len(lines))
        target = next((s for s in reversed(lines[:table_start]) if s), "")
        table_lines = lines[table_start:table_end]
        page_count = 0
        for route in ("Traditional", "Predicted"):
            try:
                route_i = table_lines.index(route)
            except ValueError:
                page_audit.append({"page": page_index + 1, "target": target,
                                   "route": route, "status": "missing_route"})
                continue
            route_lines = table_lines[route_i + 1:]
            next_route = next((j for j, s in enumerate(route_lines)
                               if s in ("Traditional", "Predicted")), None)
            if next_route is not None:
                route_lines = route_lines[:next_route]
            first_temp = next((j for j, s in enumerate(route_lines)
                               if TEMP_RE.fullmatch(s)), None)
            if first_temp is None:
                page_audit.append({"page": page_index + 1, "target": target,
                                   "route": route, "status": "missing_temperature_rows"})
                continue
            precursor_text = " ".join(s for s in route_lines[:first_temp] if s)
            precursors = re.sub(r"\s+", " ", precursor_text).strip().rstrip(",")
            for j, s in enumerate(route_lines):
                if not TEMP_RE.fullmatch(s):
                    continue
                values = []
                for raw in route_lines[j + 1:]:
                    if FLOAT_RE.fullmatch(raw):
                        values.append(float(raw))
                        if len(values) == 3:
                            break
                    elif values:
                        break
                if len(values) != 3:
                    page_audit.append({"page": page_index + 1, "target": target,
                                       "route": route, "temperature_C": int(s),
                                       "status": "incomplete_numeric_row"})
                    continue
                target_intensity, residual_intensity, phase_fraction = values
                ratio = (target_intensity / (target_intensity + residual_intensity)
                         if target_intensity + residual_intensity > 0 else 0.0)
                if abs(ratio - phase_fraction) > 0.025:
                    page_audit.append({"page": page_index + 1, "target": target,
                                       "route": route, "temperature_C": int(s),
                                       "status": "fraction_intensity_mismatch",
                                       "ratio_from_intensities": ratio,
                                       "reported_fraction": phase_fraction})
                rows.append({
                    "target_formula": target,
                    "precursor_strategy": route,
                    "precursors_as_reported": precursors,
                    "calcination_temperature_C": int(s),
                    "target_intensity_e6": target_intensity,
                    "residual_intensity_e6": residual_intensity,
                    "target_phase_fraction_reported": phase_fraction,
                    "target_phase_fraction_from_rounded_intensities": ratio,
                    "source_pdf_page_1indexed": page_index + 1,
                    "source_table": "SI3.5/Table 1",
                })
                page_count += 1
        page_audit.append({"page": page_index + 1, "target": target,
                           "rows_extracted": page_count, "status": "parsed"})
    frame = pd.DataFrame(rows)
    counts = Counter(frame["precursor_strategy"].tolist()) if not frame.empty else Counter()
    summary = {
        "status": "partial_transcription_only_stop_before_confirmatory_analysis",
        "source_pdf": pdf_path.name,
        "source_pdf_sha256": sha256(pdf_path),
        "pdf_page_count": len(doc),
        "target_table_pages_expected": 35,
        "target_table_pages_parsed": int(sum(x.get("status") == "parsed" for x in page_audit)),
        "rows_extracted": int(len(frame)),
        "unique_targets": int(frame["target_formula"].nunique()) if not frame.empty else 0,
        "rows_by_precursor_strategy": dict(counts),
        "reported_phase_fraction_rounding_check_max_abs_difference": float(
            (frame["target_phase_fraction_reported"] -
             frame["target_phase_fraction_from_rounded_intensities"]).abs().max()
        ) if not frame.empty else None,
        "article_reported_synthesis_reactions": 224,
        "unreconciled_difference_vs_article_total": 224 - int(len(frame)),
        "source_workbook_scope": "Source Data Fig. 4 workbook is aggregate Fig. 4c,d data (35 target-level rows), not the full 224 reaction-level raw XRD payload.",
        "known_stop_reasons": [
            "SI3.5 table transcription yields 202 rows, while the article reports 224 synthesis reactions; the 22-row difference is unresolved.",
            "The publisher's Supplementary Data 2 raw XRD workbook was not available locally at extraction time.",
            "No serialized V169 B0/process-aware fitted model object is present in this package, so a frozen PATH-Mat score cannot be generated.",
            "The source reports a physical target-phase fraction endpoint, which is not semantically identical to PATH-Mat's literature-derived impurity label."
        ],
        "page_audit": page_audit,
        "extraction_timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "analysis_performed": False,
    }
    return frame, summary


def audit_figure4_workbook(xlsx_path: Path) -> dict:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    result = {"file": xlsx_path.name, "sha256": sha256(xlsx_path), "sheets": {}}
    for name in wb.sheetnames:
        ws = wb[name]
        result["sheets"][name] = {"rows_including_header": ws.max_row,
                                  "columns": ws.max_column,
                                  "record_count": max(0, ws.max_row - 1)}
    result["scope_note"] = "Per publisher article page, this is Source Data Fig. 4 statistical source data, not Supplementary Data 2 raw XRD for all reactions."
    return result


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--pdf", type=Path, required=True)
    p.add_argument("--source-workbook", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True)
    a = p.parse_args()
    frame, report = extract(a.pdf)
    report["source_data_fig4_workbook_audit"] = audit_figure4_workbook(a.source_workbook)
    a.out.mkdir(parents=True, exist_ok=True)
    frame.to_csv(a.out / "ASTRAL_SI3_5_partial_phase_table_202.csv", index=False)
    (a.out / "ASTRAL_SI3_5_extraction_audit.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in ["status", "rows_extracted", "unique_targets",
          "rows_by_precursor_strategy", "unreconciled_difference_vs_article_total",
          "analysis_performed"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
