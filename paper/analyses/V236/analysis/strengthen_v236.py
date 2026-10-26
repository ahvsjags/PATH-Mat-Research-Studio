#!/usr/bin/env python3
"""Build V236's stronger, source-recomputed order-transfer account.

V236 carries forward the V235 documents and improves the evidence presentation
using source-transcribed DACHS and ZIF-62 results. It does not refit PATH-Mat,
does not change figures, and does not turn finite design points into inferential
replicates.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PARENT = ROOT.parent
SRC = PARENT / "PATH-Mat_Evidence_Strengthening_V235" / "submission_final"
OUT = ROOT / "submission_final"
RES = ROOT / "results"
REPORTS = ROOT / "reports"
SRC_RES = PARENT / "PATH-Mat_Evidence_Strengthening_V235" / "results"

MAIN_SRC = SRC / "PATH_Mat_Manuscript_Submission_V235.md"
MAIN_DOCX_SRC = SRC / "PATH_Mat_Manuscript_Submission_V235.docx"
SI_SRC = SRC / "PATH_Mat_Supporting_Information_Submission_V235.md"
SI_DOCX_SRC = SRC / "PATH_Mat_Supporting_Information_Submission_V235.docx"
MAIN_MD = OUT / "PATH_Mat_Manuscript_Submission_V236.md"
MAIN_DOCX = OUT / "PATH_Mat_Manuscript_Submission_V236.docx"
SI_MD = OUT / "PATH_Mat_Supporting_Information_Submission_V236.md"
SI_DOCX = OUT / "PATH_Mat_Supporting_Information_Submission_V236.docx"

ABSTRACT = (
    "Literature-derived models learn material chemistry and reporting structure. "
    "We evaluated PATH-Mat on 70,636 DOI- and target-formula-disjoint routes. "
    "Action-list length explained most within-source gain; generic order added no resolved increment. "
    "In 1,276 MOF diffraction reactions, source-adapted conditions ranked 492 within-formulation discordant pairs with 0.951 concordance versus 0.500 for composition. "
    "Frozen PATH-Mat correctly ranked all four temperature–phase contrasts across two unseen oxy-selenide formulas, but larger-archive transfer was heterogeneous. "
    "In AutoMOFs 4→5, linker-first selected by yield in campaign 4 beat metal-first and simultaneous addition at all five matched speeds in campaign 5 (+0.150 and +0.127 yield fraction) and lowered SAXS diameter/width by 93.6/44.9 and 105.5/51.0 nm, respectively. "
    "The same rule ranked all eight zinc-last versus zinc-earlier ZIF-62 phase contrasts correctly. "
    "Order can control outcomes and transfer across tested ZIF campaigns, but its direction depends on chemistry and endpoint; PATH-Mat does not encode reagent identity by position."
)

RESULTS_RULE_OLD = (
    "We then tested whether the direction selected by one campaign carried to the other. In a post hoc leave-one-campaign-out analysis, linker-first was selected from the mean-yield winner in AutoMOFs 4 and again when the training campaign was reversed. In the held-out AutoMOFs 5 campaign, linker-first exceeded metal-first at all five injection speeds (mean yield difference, +0.15036; range, +0.13250 to +0.17750) and had smaller mean diameter and distribution-width summaries at all five speeds. In the reverse direction, it exceeded metal-first in all four valid yield pairs and had smaller size summaries in all five matched pairs. The held-out direction therefore replicated across both campaigns and across three measured endpoints, with the largest gain in campaign 5. These are unreplicated design points in two previously inspected campaigns, so the result is a strong protocol-specific order-selection signal rather than an independent prospective policy trial (Supplementary Tables S39, S60 and S76)."
)
RESULTS_RULE_NEW = (
    "We then tested whether the order selected in one campaign beat both observed alternatives in the other. In the post hoc AutoMOFs 4→5 replay, linker-first was selected by mean yield in campaign 4 and applied unchanged to campaign 5. It exceeded metal-first and simultaneous addition at all five matched speeds (mean yield differences, +0.15036 and +0.12706, respectively). At every speed it also gave smaller SAXS mean diameter and narrower distribution width than both alternatives: L−M, −93.64 and −44.87 nm; L−T, −105.51 and −51.01 nm. Reversing the split preserved the direction: L exceeded M in all four valid yield pairs and T in all five, while diameter and width were lower than both arms at all five speeds (L−M, −51.10/−7.43 nm; L−T, −63.80/−11.50 nm). Thus, the yield-selected choice had zero observed yield regret among the three tested arms at all five held-out campaign-5 design points, and the same direction held in the reverse split wherever source outcomes were available. This is a positive cross-campaign, multi-endpoint policy-transfer result for one ZIF-8 procedure; the campaigns had been inspected and each order×speed cell contains one sample, so speed points are design conditions rather than independent replicates (Supplementary Tables S39, S60, S76 and S104)."
)

DISCUSSION_OLD = (
    "The intervention studies make synthesis order a central materials result, not merely an encoded feature. "
    "In AutoMOFs 4, linker-first was the mean-yield winner; applying that choice unchanged to AutoMOFs 5 produced higher yield at all five matched speeds (+0.15036 fraction), smaller SAXS mean diameter (−93.64 nm) and narrower distribution width (−44.87 nm). "
    "Reversing the campaign split retained the same direction at every valid design point. "
    "The ZIF-8-derived organic-before-metal rule also ranked every zinc-last ZIF-62 protocol below every zinc-earlier protocol for ZIF-zni fraction (8/8 pairwise comparisons; 7.85 versus 42.50 wt% on average). "
    "Together with controlled sequence studies in birnessite and Cu2O, order-by-context reversals in gold nanoparticles, and matched preparation-protocol contrasts in supported catalysts, these results show that route order can change phase, yield, morphology and size in measured systems. "
    "They establish cross-campaign transport of a selected order rule within ZIF-8 and directional transfer to a related framework; the Au nanoparticle reversals show why the rule is chemistry- and context-dependent (Supplementary Table S104). "
    "This is a strong experimental basis for order-aware synthesis optimization. It is not yet a PATH-Mat-generated policy: B2 discards reagent identity at each position and therefore assigns the six ZIF-62 permutations one tied score. "
    "The immediate model-development target is reagent-resolved sequence features, followed by a frozen, prospective test of model-selected order changes."
)
DISCUSSION_NEW = (
    "The strongest order result is a yield-selected rule that transported across two automation-assisted ZIF-8 campaigns and beat both measured alternatives. "
    "Selected by mean yield in AutoMOFs 4, linker-first beat metal-first and simultaneous addition at all five matched speeds in AutoMOFs 5, with mean yield gains of +0.15036 and +0.12706; it also reduced SAXS diameter and width versus both arms at every speed. "
    "The reverse campaign split retained the same direction across the available yield, diameter and width contrasts. "
    "A separate six-protocol ZIF-62 experiment supplies a chemically related but independent phase test: mean ZIF-zni fraction decreased monotonically from zinc-first (55.15 wt%) to zinc-middle (29.85 wt%) to zinc-last (7.85 wt%), and the ZIF-8-derived organic-before-metal rule ranked all eight zinc-last versus zinc-earlier pairs in the lower-impurity direction. "
    "Together, these experiments show that changing sequence can shift measured material outcomes and that an empirically selected order can transfer across tested campaigns and a related MOF. "
    "The Au nanoparticle reversals and other intervention studies show why the favorable order depends on formulation, process regime and endpoint (Supplementary Table S104). "
    "This supports order-aware synthesis selection as a practical strategy in defined systems. It does not mean the present PATH-Mat model made these choices: B2 omits reagent identity at each position and ties the six ZIF-62 permutations. "
    "The next model-level step is to represent reagent-resolved addition edges and evaluate a locked policy on a held-out intervention campaign."
)
CODE_AVAILABILITY_NEW = (
    "The analysis bundle contains frozen primary-test predictions and partition manifests; feature-ablation, nearest-route and bootstrap scripts; frozen external predictions; grouped and temporal evaluations; and source snapshots for the DACHS, ARROWS3, UiO-66, BiFeO₃, rare-earth MOF and nanoparticle analyses. V234 adds the run-matched acetic-acid analysis, and V236 adds the cross-campaign order-transfer audit with checksum-bound CSV and JSON results. A public accession for the derived PATH-Mat analysis package has not yet been assigned, and source-specific licences remain to be resolved before deposit. Raw prediction vectors for the historical reverse-direction cross-archive refit were not retained; that analysis remains a supplementary diagnostic."
)

TABLE_S104_MD = r"""
**Supplementary Table S104 | Cross-study evidence that order is a manipulable, transferable but chemistry-dependent synthesis variable.**

This evidence map consolidates controlled-order analyses in Supplementary Tables S38–S40, S60, S66–S71 and S83–S92. The strongest result is quantified separately for each baseline and endpoint; studies with different outcomes and units are not pooled.

| Evidence set | Controlled contrast / selection | Measured result | Inference supported and boundary |
|---|---|---|---|
| ZIF-8, AutoMOFs 4 → 5 [54–56] | Linker-first (L) selected as the mean-yield winner in campaign 4; applied unchanged to campaign 5 at five matched injection speeds | L−M/L−T yield +0.15036/+0.12706 (5/5 each); diameter −93.64/−105.51 nm and width −44.87/−51.01 nm (5/5 each) | Held-out campaign-5 policy-transfer result across three measured endpoints; L had zero observed yield regret among L/M/T at all five tested speeds. One sample per order×speed; campaigns previously inspected. |
| ZIF-8, reverse campaign split [54–56] | L selected by mean yield in campaign 5; applied unchanged to campaign 4 | L−M/L−T yield +0.05697 (4/4 valid)/+0.06286 (5/5); diameter −51.10/−63.80 nm and width −7.43/−11.50 nm (5/5 each) | Reverse-split directional replication across available outcomes; one source yield missing and no independent experimental replicates. |
| ZIF-62, six precursor permutations [50] | Full factorial ordering of bIm, Im and zinc nitrate; transfer the coarse organic-before-metal rule to zinc-last | Mean ZIF-zni: 7.85 wt% zinc-last, 29.85 wt% zinc-middle and 55.15 wt% zinc-first; all 12 cross-position contrasts increase with earlier zinc; zinc-last lower in all 8 zinc-last-vs-earlier pairs. bIm-before-Im is lower at all 3 fixed zinc positions. | A complete finite permutation test of order-sensitive phase composition and a directional cross-framework rule check. Replicate count and meaning of source ± values are unspecified. |
| Mn birnessite [74] | At fixed reagent ratio, concentration, addition schedule and ageing, vary last-added solution; follow-up changes oxygen exposure before oxidant addition | Oxidant-last produced detectable hausmannite/trace feitknechtite; deoxygenating the early Mn(OH)₂ intermediate removed its hausmannite signal under oxidant-last conditions | Mechanistic intervention evidence for a protocol-specific sequence-mediated phase change; source phase calls are qualitative and replicate-level uncertainty is unavailable. |
| Cu₂O nanocrystals [64] | Exchange order across chloride, nitrate and sulfate precursors | CuCl₂-before-SDS: +15% yield; morphology response to NH₂OH·HCl/NaOH reversal depended on precursor identity | Supports sequence×precursor interaction rather than a universal preferred order; source does not report batch-level uncertainty for these contrasts. |
| Co–Ru–Mg/γ-Al₂O₃ [77] | Matched OoA2/OoA3 contrasts across 24 Mg×Ru formulations; same final loading and 120 °C inter-step drying | Co-size: −2.02 nm before and −3.37 nm after ageing; growth difference −1.35 nm. Cluster-bootstrap 95% CIs: [−2.56,−1.48], [−4.27,−2.47] and [−2.21,−0.49] nm, respectively. | A coherent matched preparation-protocol effect across a factorial grid. Intervals describe design-grid consistency, not independent synthesis replication; order remains coupled to the full preparation sequence. |
| Au nanoparticles [63,67,73] | Compare order at fixed final composition across concentration, feed-rate and last-added-reagent contexts | TEM order contrast reverses (+1 to −8 nm across citrate:Au ratio; +1.7 to −3.1 nm across feed regimes); a separate fixed-composition study reports 10.0 nm when HAuCl₄ is last versus >50 nm when ethanol is last | Independent colloid studies show large, context-dependent responses; particle-level SDs do not provide independent batch uncertainty, and the preferred direction is not invariant. |

The ZIF-8 campaign-held-out comparison provides the clearest rule-transfer result because the same yield-selected choice beat both observed alternatives in the second campaign and improved yield and two size summaries. The ZIF-62 permutation supplies an independent phase endpoint and the same coarse organic-before-metal direction in a related framework. Other studies broaden the chemistry and include mechanistic evidence, but differ in replication, assignment and treatment definition. The evidence establishes order sensitivity and local rule transport in these tested systems; it does not estimate a universal success rate. Nor does it show a PATH-Mat-generated recommendation: current B2 ties reagent-resolved ZIF-62 permutations because it lacks reagent identity at each position (Supplementary Table S77). Reproduction: `analysis/strengthen_v236.py`; the source-derived contrasts and provenance are recorded in `results/V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv`, `results/V236_ZIF62_FACTORIAL_SUMMARY.csv` and `results/V236_ORDER_TRANSFER_AUDIT.json`.
""".strip()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def media_hashes(docx_path: Path) -> dict[str, str]:
    with zipfile.ZipFile(docx_path) as z:
        return {n: sha256_bytes(z.read(n)) for n in z.namelist() if n.startswith("word/media/")}


def parse_md_table(text: str, title: str, next_anchor: str) -> list[list[str]]:
    block = text.split(title, 1)[1].split(next_anchor, 1)[0]
    rows = []
    for line in block.splitlines():
        line = line.strip()
        if line.startswith("|") and not set(line.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
            cells = [c.strip() for c in line.strip("|").split("|")]
            if len(cells) == 4 and not all(set(c) <= {"-", ":"} for c in cells):
                rows.append(cells)
    return rows


def build_summaries() -> dict:
    src_csv = SRC_RES / "V235_ZIF8_ORDER_CONTRASTS.csv"
    rows = list(csv.DictReader(src_csv.open(encoding="utf-8-sig")))
    summary_rows = []
    for campaign in ("4", "5"):
        cr = [r for r in rows if r["campaign"] == campaign]
        for metric in ("yield", "diameter", "width"):
            for comparator in ("M", "T"):
                key = {
                    ("yield", "M"): "yield_L_minus_M",
                    ("yield", "T"): "yield_L_minus_T",
                    ("diameter", "M"): "diameter_L_minus_M_nm",
                    ("diameter", "T"): "diameter_L_minus_T_nm",
                    ("width", "M"): "width_L_minus_M_nm",
                    ("width", "T"): "width_L_minus_T_nm",
                }[(metric, comparator)]
                vals = [float(r[key]) for r in cr if r[key].strip()]
                summary_rows.append({
                    "campaign": campaign,
                    "contrast": f"L_minus_{comparator}",
                    "endpoint": metric,
                    "n_matched_design_points": len(vals),
                    "n_positive": sum(v > 0 for v in vals),
                    "n_negative": sum(v < 0 for v in vals),
                    "mean_difference": f"{sum(vals)/len(vals):.8f}",
                    "median_difference": f"{sorted(vals)[len(vals)//2] if len(vals)%2 else (sorted(vals)[len(vals)//2-1]+sorted(vals)[len(vals)//2])/2:.8f}",
                    "minimum": f"{min(vals):.8f}",
                    "maximum": f"{max(vals):.8f}",
                    "unit": "yield_fraction" if metric == "yield" else "nm",
                })
    with (RES / "V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(summary_rows[0]))
        w.writeheader(); w.writerows(summary_rows)

    zif62src = SRC_RES / "V235_ZIF62_PERMUTATIONS.csv"
    p_rows = list(csv.DictReader(zif62src.open(encoding="utf-8-sig")))
    pos_summary = []
    for pos in ("First", "Middle", "Last"):
        vals = [float(r["zni_wt_pct"]) for r in p_rows if r["zinc_position"] == pos]
        pos_summary.append({"zinc_position": pos, "n_protocols": len(vals), "mean_zni_wt_pct": f"{sum(vals)/len(vals):.4f}", "protocol_values_wt_pct": ";".join(f"{v:g}" for v in vals)})
    with (RES / "V236_ZIF62_FACTORIAL_SUMMARY.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(pos_summary[0]))
        w.writeheader(); w.writerows(pos_summary)
    return {
        "source_zif8_csv_sha256": hashlib.sha256(src_csv.read_bytes()).hexdigest(),
        "source_zif62_csv_sha256": hashlib.sha256(zif62src.read_bytes()).hexdigest(),
        "zif8_campaign_summary": summary_rows,
        "zif62_position_means": pos_summary,
        "design_limit": "Each order×speed cell has one sample; speeds are designed conditions, not independent synthesis replicates. No p-values or confidence intervals are inferred from these points.",
    }


def update_docs() -> dict:
    for source, destination in ((MAIN_SRC, MAIN_MD), (MAIN_DOCX_SRC, MAIN_DOCX), (SI_SRC, SI_MD), (SI_DOCX_SRC, SI_DOCX)):
        shutil.copy2(source, destination)

    main = MAIN_MD.read_text(encoding="utf-8")
    pre, rest = main.split("## Abstract\n\n", 1)
    _, tail = rest.split("\n## Main", 1)
    main = pre + "## Abstract\n\n" + ABSTRACT + "\n## Main" + tail
    if main.count(RESULTS_RULE_OLD) != 1:
        raise RuntimeError(f"Expected one Results rule-transfer paragraph, found {main.count(RESULTS_RULE_OLD)}")
    main = main.replace(RESULTS_RULE_OLD, RESULTS_RULE_NEW)
    if main.count(DISCUSSION_OLD) != 1:
        raise RuntimeError(f"Expected one Discussion synthesis paragraph, found {main.count(DISCUSSION_OLD)}")
    main = main.replace(DISCUSSION_OLD, DISCUSSION_NEW)
    code_start = main.find("The local bundle includes the frozen primary-test predictions")
    if code_start >= 0:
        code_end = main.find("\n\n## References", code_start)
        if code_end < 0:
            raise RuntimeError("Could not locate end of Code availability paragraph")
        main = main[:code_start] + CODE_AVAILABILITY_NEW + main[code_end:]
    main = main.replace("analysis/audit_and_build_v235.py", "analysis/strengthen_v236.py")
    main = main.replace("V235 adds `analysis/strengthen_v236.py`", "V236 adds `analysis/strengthen_v236.py`")
    main = main.replace("V235 adds analysis/audit_and_build_v235.py", "V236 adds analysis/strengthen_v236.py")
    main = main.replace("V235 adds `analysis/audit_and_build_v235.py`", "V236 adds `analysis/strengthen_v236.py`")
    MAIN_MD.write_text(main, encoding="utf-8")

    si = SI_MD.read_text(encoding="utf-8")
    title = "**Supplementary Table S104 |"
    if si.count(title) != 1:
        raise RuntimeError(f"Expected one S104 title, found {si.count(title)}")
    start = si.index(title)
    end = si.find("\n\n## ", start)
    if end < 0:
        raise RuntimeError("Could not find end of Table S104 section")
    si = si[:start] + TABLE_S104_MD + si[end:]
    si = si.replace("analysis/audit_and_build_v235.py", "analysis/strengthen_v236.py")
    si = si.replace("V235 cross-campaign/order-transfer arithmetic audit", "V236 cross-campaign/order-transfer audit")
    si = si.replace("V235_ORDER_TRANSFER_AUDIT.json", "V236_ORDER_TRANSFER_AUDIT.json")
    si = si.replace("V235_ZIF8_ORDER_CONTRASTS.csv", "V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv")
    si = si.replace("V235_ZIF62_PERMUTATIONS.csv", "V236_ZIF62_FACTORIAL_SUMMARY.csv")
    audit_sentence = "The V236 cross-campaign/order-transfer audit is reproduced by `analysis/strengthen_v236.py`; input and result digests are recorded in `results/V236_ORDER_TRANSFER_AUDIT.json`. "
    si = si.replace(audit_sentence * 5, audit_sentence)
    SI_MD.write_text(si, encoding="utf-8")

    m_before = media_hashes(MAIN_DOCX)
    d = Document(MAIN_DOCX)
    replaced_abstract = replaced_result = replaced_discussion = False
    for p in d.paragraphs:
        if p.text.startswith("Literature-derived models learn material chemistry and reporting practice."):
            p.text = ABSTRACT; replaced_abstract = True
        if p.text == RESULTS_RULE_OLD:
            p.text = RESULTS_RULE_NEW; replaced_result = True
        if p.text == DISCUSSION_OLD:
            p.text = DISCUSSION_NEW; replaced_discussion = True
        if "analysis/audit_and_build_v235.py" in p.text:
            p.text = p.text.replace("analysis/audit_and_build_v235.py", "analysis/strengthen_v236.py")
        if p.text.startswith("The local bundle includes the frozen primary-test predictions"):
            p.text = CODE_AVAILABILITY_NEW
    if not (replaced_abstract and replaced_result and replaced_discussion):
        raise RuntimeError(f"DOCX paragraph replacement failure: abstract={replaced_abstract}, result={replaced_result}, discussion={replaced_discussion}")
    d.save(MAIN_DOCX)
    if m_before != media_hashes(MAIN_DOCX):
        raise RuntimeError("Main DOCX media bytes changed")

    si_before = media_hashes(SI_DOCX)
    d = Document(SI_DOCX)
    # S104 is the unique four-column table whose first body row names AutoMOFs 4→5.
    table = next((t for t in d.tables if len(t.rows) == 8 and len(t.columns) == 4 and t.cell(1, 0).text.startswith("ZIF-8, AutoMOFs 4")), None)
    if table is None:
        raise RuntimeError("Could not locate the existing S104 Word table")
    for p in d.paragraphs:
        if p.text.startswith("Supplementary Table S104 | Cross-study evidence that order is a manipulable"):
            p.text = "Supplementary Table S104 | Cross-study evidence that order is a manipulable, transferable but chemistry-dependent synthesis variable."
    word_rows = [
        ["Evidence set", "Controlled contrast / selection", "Measured result", "Inference and boundary"],
        ["ZIF-8, AutoMOFs 4 → 5", "Select L by campaign-4 mean yield; apply unchanged at five matched speeds", "L−M/L−T yield +0.15036/+0.12706 (5/5 each); diameter −93.64/−105.51 nm and width −44.87/−51.01 nm (5/5 each)", "Held-out campaign-5 policy transfer; zero observed yield regret among L/M/T at all five tested speeds; one sample per order×speed"],
        ["ZIF-8 reverse split", "Select L by campaign-5 mean yield; apply to campaign 4", "Yield vs M/T +0.05697 (4/4 valid)/+0.06286 (5/5); diameter −51.10/−63.80 nm and width −7.43/−11.50 nm (5/5 each)", "Direction retained in reverse split; one source yield missing; no independent replication"],
        ["ZIF-62 six permutations", "Full bIm/Im/Zn order factorial; map organic-before-metal to zinc-last", "Zni: 7.85/29.85/55.15 wt% for zinc last/middle/first; 12/12 cross-position contrasts; 8/8 last-vs-earlier pairs", "Complete finite permutation screen; replicate counts and ± definition unavailable"],
        ["Mn birnessite", "Vary last-added solution at fixed ratio, concentrations, schedule and ageing; oxygen follow-up", "Oxidant-last: hausmannite detected; deoxygenation before oxidant removed its XRD signal", "Mechanistic sequence evidence; qualitative phase calls"],
        ["Cu2O nanocrystals", "Exchange order across chloride, nitrate and sulfate precursors", "CuCl2-before-SDS: +15% yield; morphology response to base/reductant order depended on precursor", "Sequence×precursor interaction; batch uncertainty not reported"],
        ["Co–Ru–Mg/γ-Al2O3", "Matched OoA2/OoA3 across 24 Mg×Ru formulations", "Co size: −2.02 nm pre-ageing, −3.37 nm post-ageing; growth −1.35 nm; cluster CIs exclude zero", "Coherent design-grid effect; not independent synthesis replication"],
        ["Au nanoparticle studies", "Order contrasts at matched final composition across concentration/feed/last-added contexts", "TEM differences reverse sign (+1 to −8 nm; +1.7 to −3.1 nm); separate last-addition contrast 10 nm vs >50 nm", "Large but context-dependent response; particle SD is not batch uncertainty"],
    ]
    for i, row in enumerate(word_rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    for p in d.paragraphs:
        if "The strongest rule-transfer result is a campaign-held-out ZIF-8 comparison" in p.text:
            p.text = "The strongest rule-transfer result is the three-arm campaign-held-out ZIF-8 comparison; other systems broaden the chemistry and mechanisms but differ in randomization and replicate reporting. These interventions establish order sensitivity in defined workflows, not a universal order policy or a PATH-Mat-generated recommendation."
        if "V235 cross-campaign/order-transfer arithmetic audit" in p.text:
            p.text = p.text.replace("V235 cross-campaign/order-transfer arithmetic audit", "V236 cross-campaign/order-transfer audit")
        if "analysis/audit_and_build_v235.py" in p.text:
            p.text = p.text.replace("analysis/audit_and_build_v235.py", "analysis/strengthen_v236.py")
        if p.text.startswith("The local bundle includes the frozen primary-test predictions"):
            p.text = CODE_AVAILABILITY_NEW
        audit_sentence = "The V236 cross-campaign/order-transfer audit is reproduced by analysis/strengthen_v236.py; input and result digests are recorded in results/V236_ORDER_TRANSFER_AUDIT.json."
        if audit_sentence in p.text:
            first = p.text.find(audit_sentence)
            after = p.text[first + len(audit_sentence):].replace(audit_sentence, "")
            p.text = p.text[:first] + audit_sentence + after
    d.save(SI_DOCX)
    if si_before != media_hashes(SI_DOCX):
        raise RuntimeError("SI DOCX media bytes changed")
    return {"main_media": m_before, "si_media": si_before}


def remove_parent_version_copies() -> None:
    """Keep the V236 deliverable directory unambiguous; parent V235 stays intact."""
    stale = [
        ROOT / "submission_final" / "PATH_Mat_Manuscript_Submission_V235.md",
        ROOT / "submission_final" / "PATH_Mat_Manuscript_Submission_V235.docx",
        ROOT / "submission_final" / "PATH_Mat_Supporting_Information_Submission_V235.md",
        ROOT / "submission_final" / "PATH_Mat_Supporting_Information_Submission_V235.docx",
        ROOT / "analysis" / "audit_and_build_v235.py",
        ROOT / "results" / "V235_ORDER_TRANSFER_AUDIT.json",
        ROOT / "results" / "V235_ZIF62_PERMUTATIONS.csv",
        ROOT / "results" / "V235_ZIF8_ORDER_CONTRASTS.csv",
        ROOT / "reports" / "V235_ORDER_AND_TRANSFER_STRENGTHENING_REPORT.md",
    ]
    for path in stale:
        resolved = path.resolve()
        resolved.relative_to(ROOT.resolve())
        if resolved.is_file():
            resolved.unlink()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True); RES.mkdir(parents=True, exist_ok=True); REPORTS.mkdir(parents=True, exist_ok=True)
    stats = build_summaries()
    media = update_docs()
    main_txt = MAIN_MD.read_text(encoding="utf-8")
    abstract = main_txt.split("## Abstract\n\n", 1)[1].split("\n## Main", 1)[0]
    word_count = len(re.findall(r"\b[\w’−+.]+(?:[-–][\w’−+.]+)*\b", abstract))
    si_txt = SI_MD.read_text(encoding="utf-8")
    if word_count > 150:
        raise RuntimeError(f"Abstract exceeds 150 words: {word_count}")
    if si_txt.count("**Supplementary Table S104 |") != 1:
        raise RuntimeError("S104 Markdown title count is not one")
    v235 = json.loads((SRC_RES / "V235_ORDER_TRANSFER_AUDIT.json").read_text(encoding="utf-8"))
    audit = {
        "schema": "pathmat_v236_order_transfer_reanalysis",
        "parent_audit": "V235_ORDER_TRANSFER_AUDIT.json",
        "parent_audit_sha256": hashlib.sha256((SRC_RES / "V235_ORDER_TRANSFER_AUDIT.json").read_bytes()).hexdigest(),
        "output_sha256": {
            "V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv": hashlib.sha256((RES / "V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv").read_bytes()).hexdigest(),
            "V236_ZIF62_FACTORIAL_SUMMARY.csv": hashlib.sha256((RES / "V236_ZIF62_FACTORIAL_SUMMARY.csv").read_bytes()).hexdigest(),
        },
        "source_order_rule": v235["order_rule"],
        "recomputed_summary": stats,
        "interpretation": {
            "positive": ["yield-selected linker-first order beats both recorded alternatives at every held-out AutoMOFs 5 speed point", "reverse split preserves every available comparison direction", "ZIF-62 six-permutation data show a monotonic measured phase gradient and the same coarse organic-before-metal direction"],
            "scope": "specific cross-campaign rule transport and finite controlled protocol contrasts in related ZIF workflows; not universal zero-shot transfer, a universal optimal order, or a PATH-Mat-generated causal policy",
            "uncertainty": "No inferential p-value or confidence interval is derived from unreplicated speed-design points or the six ZIF-62 protocols."
        },
        "abstract_word_count": word_count,
        "figure_assets_unchanged": True,
        "docx_media_sha256": media,
    }
    (RES / "V236_ORDER_TRANSFER_AUDIT.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    report = f"""# PATH-Mat V236 — strengthened order-transfer evidence

## New positive result

The manuscript now reports the complete matched three-arm comparison, rather than only L−M. A linker-first (L) route selected by mean yield in AutoMOFs 4 beat both metal-first and simultaneous addition in AutoMOFs 5 at all five matched speed design points. Mean held-out differences were +0.15036 and +0.12706 yield fraction; SAXS diameter and width were lower than both alternatives at every speed. Reversing the campaign split preserved the direction in every available comparison. Thus, the observed held-out route choice had zero yield regret among the three tested arms at all five AutoMOFs 5 speeds, for these campaigns and recorded conditions.

The independent six-permutation ZIF-62 screen showed a monotonic measured ZIF-zni gradient from 55.15 wt% (zinc first) through 29.85 wt% (middle) to 7.85 wt% (last); the ZIF-8-derived organic-before-metal direction ranked all eight last-versus-earlier comparisons consistently. The finite factorial also showed that linker sub-order matters at each zinc position.

## Strongest supported wording

- **Zero-shot:** PATH-Mat has a successful, bounded formula-held-out result (four correct temperature–phase rankings across two formulas), alongside heterogeneous and negative transfers. The evidence supports *task-specific zero-shot success*, not a universal transfer claim.
- **Order effect:** controlled experiments show that order changes measured outcomes in specific systems. The new V236 Results quantify a three-arm, cross-campaign ZIF-8 rule transfer across yield, particle diameter and width, plus a full ZIF-62 permutation gradient. These are positive physical interventions, not merely feature correlations.
- **Optimization:** an empirically chosen order rule transfers between the two ZIF-8 campaigns and beats both observed alternative orders in the held-out campaign. Separately, the nested source-adapted rare-earth-MOF policy reached 43.8% named-product outcomes versus 40.9% for random choice (+2.86 percentage points; 95% metal-cluster CI, +1.60 to +4.81), while a training-fold-selected global 80 °C/2 M recipe reached 45.1%. These are positive offline decision-value results, but the simple global rule remains stronger and neither analysis is a prospective PATH-Mat-generated causal policy. B2 still cannot encode reagent identity at each sequence position.

## Statistical boundary

Each campaign×order×speed cell has one sample; injection speeds are designed conditions, not independent replicates. The ZIF-62 source does not define the independent replicate count or the meaning of the reported ± values. V236 therefore adds no p-values or confidence intervals to these finite comparisons. The campaigns were inspected previously, so this is a stronger reanalysis of existing public evidence, not a blinded confirmation.

## File changes and validation

- Main text: full L−M and L−T effect sizes in both campaign directions; abstract tightened to {word_count} words.
- Supporting Information: Supplementary Table S104 updated with both comparators and the exact position gradient.
- Figures: unchanged. Main DOCX media files: {len(media['main_media'])}; SI DOCX media files: {len(media['si_media'])}; SHA-256 maps preserved exactly from the V235 parent documents.
- Reproduction script: `analysis/strengthen_v236.py`.
- Machine-readable outputs: `results/V236_ZIF8_ORDER_TRANSFER_SUMMARY.csv`, `results/V236_ZIF62_FACTORIAL_SUMMARY.csv`, `results/V236_ORDER_TRANSFER_AUDIT.json`.

## Sources

- AutoMOFs 4 and 5 primary data: [Zenodo AutoMOFs 4](https://doi.org/10.5281/zenodo.11237900), [Zenodo AutoMOFs 5](https://doi.org/10.5281/zenodo.11237988), and the DACHS synthesis study ([ChemRxiv DOI](https://doi.org/10.26434/chemrxiv-2025-7fgg0)).
- ZIF-62 six-permutation experiment: [Nature Communications article](https://doi.org/10.1038/s41467-025-64583-9).
"""
    (REPORTS / "V236_ORDER_TRANSFER_STRENGTHENING_REPORT.md").write_text(report, encoding="utf-8")
    remove_parent_version_copies()
    print(json.dumps({"version": "V236", "abstract_words": word_count, "figure_media_unchanged": True, "main_media_count": len(media["main_media"]), "si_media_count": len(media["si_media"]), "new_results": ["L beats M and T at all five held-out AutoMOFs 5 speeds", "reverse split same direction", "ZIF-62 monotonic phase gradient"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
