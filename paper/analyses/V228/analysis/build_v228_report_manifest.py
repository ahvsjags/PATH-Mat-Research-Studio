from pathlib import Path
import hashlib,json,zipfile,csv
from docx import Document
ROOT=Path(r'D:\Downloads\材料agent\PATH-Mat_Evidence_Strengthening_V228')
SUB=ROOT/'submission_final'; RES=ROOT/'results'; ANA=ROOT/'analysis'
paths=[SUB/'PATH_Mat_Manuscript_Submission_V228.md',SUB/'PATH_Mat_Manuscript_Submission_V228.docx',SUB/'PATH_Mat_Supporting_Information_Submission_V228.md',SUB/'PATH_Mat_Supporting_Information_Submission_V228.docx',RES/'V228_BOLDRIN_OOA2_VS_OOA3_MATCHED_CELLS.csv',RES/'V228_BOLDRIN_MATCHED_ORDER_AUDIT.json',ANA/'reanalyze_boldrin_matched_order.py',ANA/'build_v228_submission.py',ANA/'build_v228_report_manifest.py',ROOT/'V228_EVIDENCE_STRENGTHENING_REPORT.md']
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def media(p):
 with zipfile.ZipFile(p) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
mainmd=(SUB/'PATH_Mat_Manuscript_Submission_V228.md').read_text(encoding='utf-8')
simd=(SUB/'PATH_Mat_Supporting_Information_Submission_V228.md').read_text(encoding='utf-8')
absx=mainmd.split('## Abstract\n\n')[1].split('\n## Main',1)[0]
csvrows=list(csv.DictReader((RES/'V228_BOLDRIN_OOA2_VS_OOA3_MATCHED_CELLS.csv').open(encoding='utf-8')))
audit=json.loads((RES/'V228_BOLDRIN_MATCHED_ORDER_AUDIT.json').read_text(encoding='utf-8'))
mdoc=Document(SUB/'PATH_Mat_Manuscript_Submission_V228.docx')
sidoc=Document(SUB/'PATH_Mat_Supporting_Information_Submission_V228.docx')
checks={
 'abstract_words':len(absx.split()),
 'matched_cells':len(csvrows),
 'unique_Mg_Ru_pairs':len({(r['Mg_wt_pct'],r['Ru_wt_pct']) for r in csvrows}),
 'new_SI_tables_present':('Supplementary Table S91' in simd and 'Supplementary Table S92' in simd),
 'S1_to_S92_index_present':'Supporting Tables S1–S92' in mainmd,
 'main_docx_paragraphs':len(mdoc.paragraphs),
 'SI_docx_tables':len(sidoc.tables),
 'S91_shape':[len(sidoc.tables[80].rows),len(sidoc.tables[80].columns)],
 'S92_shape':[len(sidoc.tables[81].rows),len(sidoc.tables[81].columns)],
 'no_cell_level_p_values':not any('sign_test_p' in k for d in audit['results'].values() if isinstance(d,dict) for k in d),
 'main_media_unchanged':media(SUB/'PATH_Mat_Manuscript_Submission_V227.docx')==media(SUB/'PATH_Mat_Manuscript_Submission_V228.docx'),
 'SI_media_unchanged':media(SUB/'PATH_Mat_Supporting_Information_Submission_V227.docx')==media(SUB/'PATH_Mat_Supporting_Information_Submission_V228.docx')}
assert checks['abstract_words']<=150
assert checks['matched_cells']==checks['unique_Mg_Ru_pairs']==24
assert checks['S91_shape']==[4,5] and checks['S92_shape']==[25,9]
assert checks['main_media_unchanged'] and checks['SI_media_unchanged']
report="""# V228 evidence-strengthening report

## What changed

V228 adds a source-traceable matched reanalysis of the published Co–Ru–Mg/γ-Al₂O₃ factorial study. The OoA2–OoA3 contrast holds intermediate drying and final heat treatment constant while reversing whether Mg is deposited before or after Co and Ru. OoA1, with its 550 °C intermediate calcination, is excluded from this order-specific contrast. The analysis averages the two technical XRD wells per preparation before pairing 24 matched Mg×Ru formulations.

## Main result

With Δ defined as OoA2−OoA3, the mean XRD-derived Scherrer Co-size contrast was **−2.02 nm before ageing** (95% t interval across eight Mg-loading cluster means: **−2.56 to −1.48 nm**; lower with OoA2 in all 8 Mg clusters and 23/24 Mg×Ru cells). After 100 h ageing, the contrast was **−3.37 nm** (**−4.27 to −2.47 nm**; lower in 8/8 clusters and 24/24 cells). Mean ageing growth was 0.95 nm for OoA2 versus 2.30 nm for OoA3; the matched growth contrast was **−1.35 nm** (**−2.21 to −0.49 nm**), with lower growth in 7/8 Mg clusters and 17/24 cells. The cluster-bootstrap sensitivity intervals are in the audit JSON. Omitting any one Mg-loading level left each average contrast below zero (pre-ageing −2.18 to −1.92 nm; post-ageing −3.67 to −3.18 nm; growth −1.58 to −1.14 nm).

These grid means correspond to 16.2% lower mean Scherrer size before ageing and 22.8% lower after ageing, relative to OoA3; mean ageing growth was 58.7% lower. Percentages are descriptive contrasts over this factorial grid.

## Interpretation and boundary

This materially strengthens evidence that operation order can alter a measured physical outcome in a defined materials workflow. It is an order-specific contrast with matched composition, drying schedule and final treatment, directionally consistent across the tested Mg-loading grid and accompanied by an ageing endpoint. The complete ZIF-62 permutation experiment and other controlled-order studies already reported in the manuscript provide separate, system-specific corroboration.

This is a retrospective reanalysis of one published catalyst library, not a new replication. Duplicate wells are technical measurements and were averaged; the source does not document independent synthesis replicates for each factorial cell or randomized protocol allocation. Intervals summarize consistency across the eight tested Mg levels, not a random population of materials systems. The result supports an **order-sensitive intervention effect in this defined catalyst workflow**; it does not establish a universal order rule, prove universal zero-shot generalization, show that PATH-Mat’s current generic order feature predicts the effect, or demonstrate PATH-Mat-guided causal optimization. The Lee primary B2−B1 increment remains statistically unresolved.

## Source and methods

Order definitions, inter-impregnation drying (120 °C, 4 h), final drying (120 °C, 4 h), final calcination (250 °C, 8 h), two-well design and the Scherrer Co-size endpoint were checked against the official Supporting Information. Tables S3 and S4 provide pre- and post-ageing observations. Each pair matches Mg and Ru loading. The primary interval averages the three Ru-specific contrasts within each Mg level, then calculates a t interval across the eight Mg-level means. A 100,000-draw percentile bootstrap resamples Mg levels while retaining all Ru strata (seed 20261016), and a leave-one-Mg-level-out check assesses dependence on individual grid levels. No p-value treating the 24 formulation cells or 48 technical wells as independent replicates is used.

## Files

- Main manuscript: `submission_final/PATH_Mat_Manuscript_Submission_V228.docx` and `.md`
- Supporting Information: `submission_final/PATH_Mat_Supporting_Information_Submission_V228.docx` and `.md`
- Cellwise paired data: `results/V228_BOLDRIN_OOA2_VS_OOA3_MATCHED_CELLS.csv`
- Machine-readable audit: `results/V228_BOLDRIN_MATCHED_ORDER_AUDIT.json`
- Reanalysis script: `analysis/reanalyze_boldrin_matched_order.py`
- Submission build script: `analysis/build_v228_submission.py`

Supporting Tables S91 (summary) and S92 (all 24 matched cells) contain the new analysis. Figure artwork was not edited: embedded `word/media` payloads in the main manuscript and SI are byte-identical to V227.

## Verification

"""+'\n'.join(f'- `{k}`: **{v}**' for k,v in checks.items())+"""

## Sources

- [Boldrin et al., Chemical Science (2015)](https://pubs.rsc.org/en/content/articlelanding/2015/sc/c4sc02116a)
- [Official Supporting Information, Tables S3–S4](https://www.rsc.org/suppdata/sc/c4/c4sc02116a/c4sc02116a1.pdf)
"""
(ROOT/'V228_EVIDENCE_STRENGTHENING_REPORT.md').write_text(report,encoding='utf-8')
manifest={'package':'PATH-Mat evidence strengthening V228','parent':'V227 submission package','scope':'Matched retrospective reanalysis of published OoA2 versus OoA3 catalyst XRD observations; no new synthesis or figure artwork changes.','sources':[{'citation':'Boldrin et al., Chemical Science 6, 935–944 (2015), DOI 10.1039/C4SC02116A','url':'https://pubs.rsc.org/en/content/articlelanding/2015/sc/c4sc02116a'},{'citation':'Official Supporting Information, synthesis protocol and Tables S3–S4','url':'https://www.rsc.org/suppdata/sc/c4/c4sc02116a/c4sc02116a1.pdf'}],'checks':checks,'files':{str(p.relative_to(ROOT)).replace('\\','/'): {'sha256':sha(p),'bytes':p.stat().st_size} for p in paths}}
(ROOT/'V228_ARTIFACT_MANIFEST.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(checks,indent=2,ensure_ascii=False))
